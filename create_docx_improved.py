import sys
import subprocess
try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

# Estilos
title = doc.add_heading('Guión de Exposición (Navegación GitHub) - Fase 2 EcoLogCity', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_heading('Roles de Exposición (3 Integrantes)', level=1)
doc.add_paragraph('• ORADOR 1: Albornoz Peña Jeferson Bener (Scrum Master / Analista de Negocio)\n'
                  '• ORADOR 2: Quispe Aquino Junior (Desarrollador Full-Stack y UX/UI)\n'
                  '• ORADOR 3: Landa Rojas Alexander Nelson (Arquitecto Cloud y Especialista Backend y Operador de Pantalla)')

doc.add_paragraph('Nota para Alexander (Operador de Pantalla): Mantén el zoom del navegador entre 110% y 125% para que todo sea legible desde el proyector. Tú controlarás la pantalla durante toda la presentación, siguiendo las instrucciones marcadas como [ACCIÓN EN PANTALLA].').bold = True

# ================= ORADOR 1 =================
doc.add_heading('🗣️ ORADOR 1: Albornoz Peña (Introducción y Artefacto 1)', level=2)

p = doc.add_paragraph()
p.add_run('[ACCIÓN EN PANTALLA - ALEXANDER]: Mostrar la página principal del repositorio en GitHub (el README.md). Hacer scroll suave mostrando la tabla de los 3 integrantes.\n\n').font.color.rgb = RGBColor(255, 0, 0)

p.add_run('[ORADOR 1 - ALBORNOZ]: ').bold = True
p.add_run('Buenas [tardes/noches] profesor y compañeros. Hoy nuestro equipo, conformado por Junior Quispe, Alexander Landa y mi persona, Jeferson Albornoz, presentará la Fase 2 de planificación del proyecto EcoLogCity.\n\n')
p.add_run('Para empezar, explicaré el Artefacto 1: Transformación Ágil.\n\n')

p.add_run('[ACCIÓN EN PANTALLA - ALEXANDER]: Navegar en vivo hacia la carpeta docs/02 Planificación/ y hacer clic en el Artefacto 1 (01 Transformando a ágil V_1_0_0.md).\n\n').font.color.rgb = RGBColor(255, 0, 0)

p.add_run('[ORADOR 1 - ALBORNOZ]: ').bold = True
p.add_run('¿Qué hicimos y cómo lo hicimos? Convertimos nuestros requerimientos funcionales monolíticos en Épicas e Historias de Usuario (US). Como pueden ver en pantalla, usamos la plantilla canónica de "Como [Rol], quiero [Acción], para [Beneficio]".\n\n')

p.add_run('[ACCIÓN EN PANTALLA - ALEXANDER]: Hacer scroll hasta los Criterios de Aceptación (Gherkin) y señalar/seleccionar el bloque con el mouse.\n\n').font.color.rgb = RGBColor(255, 0, 0)

p.add_run('[ORADOR 1 - ALBORNOZ]: ').bold = True
p.add_run('Además, cada historia posee criterios de aceptación bajo la sintaxis BDD Gherkin (Dado, Cuando, Entonces). ¿Para qué hicimos esto? Para asegurar que cada unidad de desarrollo sea testeable. Finalmente, establecimos un Definition of Done (DoD) estricto.\n\n')

p.add_run('[ACCIÓN EN PANTALLA - ALEXANDER]: Hacer scroll hasta el final del Artefacto 1 y resaltar las reglas del 80% de cobertura, SonarQube y Peer Review.\n\n').font.color.rgb = RGBColor(255, 0, 0)

p.add_run('[ORADOR 1 - ALBORNOZ]: ').bold = True
p.add_run('Como se observa, nuestro DoD exige un 80% de cobertura y código limpio. Ahora daré pase a mi compañero Junior para que muestre el Artefacto 2.\n')


# ================= ORADOR 2 =================
doc.add_heading('🗣️ ORADOR 2: Quispe Aquino (Artefactos Jira y Mitad de Riesgos)', level=2)

p2 = doc.add_paragraph()
p2.add_run('[ACCIÓN EN PANTALLA - ALEXANDER]: Volver al README, abrir el Artefacto 2 (Artefactos Jira) y hacer scroll lento mostrando las capturas recortadas de Jira.\n\n').font.color.rgb = RGBColor(255, 0, 0)

p2.add_run('[ORADOR 2 - JUNIOR]: ').bold = True
p2.add_run('Gracias Jeferson. Pasando al Artefacto 2, llevamos todo este planeamiento a Jira Software.\n')
p2.add_run('¿Qué y cómo lo hicimos? Primero armamos nuestro Roadmap. Luego, en el Backlog, estimamos el esfuerzo de las historias utilizando Puntos de Historia con la sucesión de Fibonacci (3, 5, 8). Además, planificamos el Sprint 1 fijando un "Sprint Goal" enfocado en el núcleo de la base de datos, y configuramos nuestro Tablero Scrum.\n')
p2.add_run('¿Para qué lo hicimos? Para tener trazabilidad total y transparencia del trabajo empírico de nuestro equipo de 3 desarrolladores.\n\n')

p2.add_run('[ACCIÓN EN PANTALLA - ALEXANDER]: Regresar a la carpeta docs/02 Planificación/ y abrir el Artefacto 3 (Registro de riesgos). Resaltar las columnas Prob. e Imp. en la tabla.\n\n').font.color.rgb = RGBColor(255, 0, 0)

p2.add_run('[ORADOR 2 - JUNIOR]: ').bold = True
p2.add_run('Respecto al Artefacto 3: Registro de Riesgos. Construimos una matriz cuantitativa según el PMBOK. Calculamos la Severidad multiplicando Probabilidad por Impacto. Ahora, mi compañero Alexander continuará explicando la mitigación y el presupuesto.\n')


# ================= ORADOR 3 =================
doc.add_heading('🗣️ ORADOR 3: Landa Rojas (Riesgos, Presupuesto y Git)', level=2)

p3 = doc.add_paragraph()
p3.add_run('[ACCIÓN EN PANTALLA - ALEXANDER]: Detenerse en la fila de la tabla de riesgos que habla del algoritmo VRP y la altitud de Huancayo. Subrayar el texto de mitigación con el mouse.\n\n').font.color.rgb = RGBColor(255, 0, 0)

p3.add_run('[ORADOR 3 - ALEXANDER]: ').bold = True
p3.add_run('Gracias Junior. ¿Para qué hicimos esta matriz? Para prevenir en lugar de reaccionar. Por ejemplo, identificamos como riesgo de "Severidad Alta (15)" el cálculo del algoritmo VRP debido a la altitud de Huancayo. Nuestro plan preventivo mitiga esto desde el diseño.\n\n')

p3.add_run('[ACCIÓN EN PANTALLA - ALEXANDER]: Volver al menú anterior y abrir el Artefacto 4 (Presupuesto del proyecto). Bajar mostrando el CAPEX y OPEX.\n\n').font.color.rgb = RGBColor(255, 0, 0)

p3.add_run('[ORADOR 3 - ALEXANDER]: ').bold = True
p3.add_run('Finalmente, elaboramos el Presupuesto del Proyecto (Artefacto 4). ¿Cómo lo hicimos? Desglosamos el costo de los 3 miembros de nuestro equipo en el CAPEX, sumamos el OPEX (licencias y nube), y lo más importante: incluimos una reserva de contingencia exacta del 12% basada en los riesgos evaluados anteriormente.\n\n')

p3.add_run('[ACCIÓN EN PANTALLA - ALEXANDER]: Hacer scroll al final de la tabla financiera y resaltar el 12%. Luego volver al README principal (usando el botón "Volver al README") y señalar en el lado derecho el Tag de Release v1.0.0.\n\n').font.color.rgb = RGBColor(255, 0, 0)

p3.add_run('[ORADOR 3 - ALEXANDER]: ').bold = True
p3.add_run('Para concluir, ¿Para qué hicimos todo en GitHub? Para evidenciar la Integración y Calidad Documental. Como se observa, el repositorio cuenta con navegación bidireccional mediante enlaces ágiles, versionamiento semántico estricto (V_1_0_0) y todo consolidado limpiamente en Markdown. Con esto concluimos nuestra sustentación, gracias.\n')

doc.save('Guion_Exposicion_Final.docx')
print("Documento mejorado creado.")
