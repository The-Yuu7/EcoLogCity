import sys
import subprocess

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

# Estilos de Título
title = doc.add_heading('Guión de Exposición: Fase 2 - Planificación Ágil', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_paragraph('Proyecto: EcoLogCity (Plataforma de Optimización de Rutas Sostenibles)\n').alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_heading('Estructura de la Presentación (3 Integrantes)', level=1)
doc.add_paragraph('Este guión está diseñado para que la transición entre los 3 expositores sea fluida. Cada uno debe explicar qué se hizo, cómo se hizo y para qué se hizo, justificándose en la rúbrica de evaluación.')

# ================= ORADOR 1 =================
doc.add_heading('🗣️ ORADOR 1: Introducción y Artefacto 1 (Transformación Ágil)', level=2)
p1 = doc.add_paragraph()
p1.add_run('[ORADOR 1]: ').bold = True
p1.add_run('Buenas [tardes/noches] profesor y compañeros. El día de hoy, nuestro equipo presentará la Fase 2 de planificación del proyecto EcoLogCity, donde hemos llevado nuestros requerimientos monolíticos iniciales hacia un entorno ágil y profesional.\n\n')
p1.add_run('Para empezar, explicaré el Artefacto 1, enfocado en la transformación a ágil.\n')
p1.add_run('¿Qué hicimos? ').bold = True
p1.add_run('Convertimos nuestros requerimientos funcionales y no funcionales en Épicas, Historias de Usuario (US) y Enablers o Historias Técnicas.\n')
p1.add_run('¿Cómo lo hicimos? ').bold = True
p1.add_run('Utilizamos la plantilla canónica de "Como [Rol], quiero [Acción], para [Beneficio]". Además, a cada historia le añadimos criterios de aceptación usando el lenguaje BDD Gherkin, estructurando escenarios con "Dado, Cuando, Entonces". Finalmente, definimos un Definition of Done (DoD) estricto que exige un 80% de cobertura de pruebas, código limpio en SonarQube y validación por Peer Review mediante Pull Requests.\n')
p1.add_run('¿Para qué lo hicimos? ').bold = True
p1.add_run('Para asegurar que cada pequeño fragmento que desarrollemos entregue valor real al usuario final, y garantizar que no daremos ninguna tarea por "terminada" hasta que cumpla estándares internacionales de calidad y seguridad.\n\n')
p1.add_run('Ahora, cederé la palabra a mi compañero/a para que muestre cómo llevamos todo esto a la práctica en Jira.')

# ================= ORADOR 2 =================
doc.add_heading('🗣️ ORADOR 2: Artefactos Jira y Mitad de Riesgos', level=2)
p2 = doc.add_paragraph()
p2.add_run('[ORADOR 2]: ').bold = True
p2.add_run('Muchas gracias. Pasando al Artefacto 2, nos centramos en la Configuración en Jira Software.\n\n')
p2.add_run('¿Qué hicimos? ').bold = True
p2.add_run('Parametrizamos un entorno empresarial de gestión ágil en Atlassian Jira.\n')
p2.add_run('¿Cómo lo hicimos? ').bold = True
p2.add_run('Primero, armamos nuestro Roadmap alineando las Épicas. Luego, en el Backlog, estimamos el esfuerzo de cada historia usando Story Points con la serie de Fibonacci (1, 2, 3, 5, 8). Planificamos nuestro Sprint 1 (de 2 semanas) con un Sprint Goal claro enfocado en sentar la arquitectura base y la geocodificación. Finalmente, organizamos nuestro Tablero Scrum y creamos el Release v1.0.0-MVP.\n')
p2.add_run('¿Para qué lo hicimos? ').bold = True
p2.add_run('Para tener total trazabilidad empírica del avance del equipo. Jira nos permite evidenciar bloqueos rápidamente y mantener una transparencia absoluta sobre qué se está construyendo.\n\n')
p2.add_run('Pasando al Artefacto 3: Registro de Riesgos, iniciamos la gestión basada en el estándar PMBOK.\n')
p2.add_run('¿Qué hicimos y cómo lo hicimos? ').bold = True
p2.add_run('Construimos una matriz cuantitativa calculando la "Severidad", multiplicando la Probabilidad (del 1 al 5) por el Impacto (del 1 al 5). Mapeamos riesgos clave de infraestructura, de capital humano y técnicos.\n\n')
p2.add_run('A continuación, mi compañero/a detallará cómo mitigamos estos riesgos y expondrá el presupuesto del proyecto.')

# ================= ORADOR 3 =================
doc.add_heading('🗣️ ORADOR 3: Cierre de Riesgos, Presupuesto y Git', level=2)
p3 = doc.add_paragraph()
p3.add_run('[ORADOR 3]: ').bold = True
p3.add_run('Gracias. Continuando con el Artefacto 3 sobre los riesgos...\n')
p3.add_run('¿Para qué lo hicimos? ').bold = True
p3.add_run('El objetivo de esta matriz no es reaccionar al caos, sino prevenirlo. Por ejemplo, identificamos como riesgo de "Severidad Alta (15)" el impacto de la altitud de Huancayo en el algoritmo VRP. Al prevenirlo, asignamos planes de mitigación preventivos desde el diseño arquitectónico y definimos roles responsables exactos en caso requiramos activar un plan de contingencia.\n\n')
p3.add_run('Con estos riesgos mapeados, elaboramos el Artefacto 4: Presupuesto del Proyecto.\n')
p3.add_run('¿Qué hicimos y cómo lo hicimos? ').bold = True
p3.add_run('Diseñamos un modelo financiero a 4 meses. Desglosamos el CAPEX (Recursos Humanos multiplicando las horas por la tarifa en USD), y el OPEX (Servidores Cloud en AWS y base de datos). A esto sumamos el costo de licencias de Jira y SonarQube. Muy importante: gracias a la matriz de riesgos previa, logramos justificar una reserva de contingencia exacta del 12% sobre el subtotal.\n')
p3.add_run('¿Para qué lo hicimos? ').bold = True
p3.add_run('Para demostrar la viabilidad económica del proyecto ante DistriRápido S.A.C., evidenciando exactamente a dónde va destinado cada dólar del desarrollo.\n\n')
p3.add_run('Finalmente, ¿Dónde y cómo entregamos este trabajo? ').bold = True
p3.add_run('Como exige la Integración Continua (Criterio 5), todo se encuentra alojado en GitHub en la ruta docs/02 Planificación/. Hemos creado enlaces bidireccionales ágiles: desde el README principal pueden llegar a cualquier artefacto, y desde cada artefacto retornar al menú. Todo bajo el formato estricto Markdown y versionamiento semántico V_1_0_0.\n\n')
p3.add_run('Con esto concluimos la sustentación de la Fase 2. Muchas gracias.')

doc.save('Guion_Exposicion_Fase2.docx')
print("Documento DOCX generado exitosamente.")
